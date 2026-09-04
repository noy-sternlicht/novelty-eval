"""Prepare one annotation batch per annotator, plus the decoding key.

Each annotator gets the same batch three ways -- an xlsx to fill in, an HTML page
to annotate in a browser, and an Apps Script that builds their Google Form for
annotating on a phone -- so they can pick whichever route suits them.

A batch is built from named groups, each drawing on its own judge run: the
mistakes under study, controls the judge got right, hard pairs a strong judge
failed, and whatever else a study needs. Groups may come from different runs
and different benchmarks.

Every pair is shown blind -- A/B order randomized, no scores, no gold label,
and the groups shuffled together with nothing marking which is which. Rows are
identified only by an opaque item_id, because problem ids collide across
benchmarks and would otherwise both clash and hint at a row's group. key.xlsx
maps those ids back to gold, and is never distributed.

Usage:
    python -m novelty_eval.judge_error_annotation.prepare_annotation_batch \
        --config src/novelty_eval/judge_error_annotation/annotation_batch.yaml
"""

from __future__ import annotations

import argparse
import random
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from novelty_eval.judge_error_annotation.annotation_page import write_annotation_page
from novelty_eval.judge_error_annotation.google_form_export import build_apps_script
from novelty_eval.judge_error_annotation.selection import (
    ScoredInstance,
    is_judge_mistake,
    judge_novelty_winner,
    novelty_margin,
    partition_scored,
)

# What the annotator fills in, and what they read. The two answer columns sit
# before the long idea columns so they stay on screen while reading.
SHEET_HEADERS = ["item_id", "area", "your_choice", "notes", "idea_A", "idea_B"]
COLUMN_WIDTHS = {"A": 10, "B": 26, "C": 13, "D": 30, "E": 70, "F": 70}
WRAPPED_COLUMNS = ("B", "D", "E", "F")
CHOICES = ["A", "B", "unclear"]

# Which side of a run's partition a group draws from.
MISTAKES, CONTROLS = "mistakes", "controls"
# How instances are spread across annotators.
DISJOINT, SHARED = "disjoint", "shared"

# Same palette as the HTML report in annotate_judge_errors.py, so the artifacts
# this module hands out look like they belong together.
ACCENT, RULE, INPUT_BG, BAND_BG = "2F6FB5", "D8DEE5", "EDF3FA", "F7F9FB"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
INPUT_FILL = PatternFill("solid", fgColor=INPUT_BG)
BAND_FILL = PatternFill("solid", fgColor=BAND_BG)
RULED = Border(bottom=Side(style="thin", color=RULE), right=Side(style="thin", color=RULE))

KEY_HEADERS = [
    "annotator",
    "item_id",
    "group",
    "source",
    "problem_id",
    "is_mistake",
    "shown_A_idea",
    "shown_B_idea",
    "gold_idea",
    "judge_novelty_winner",
]

INSTRUCTIONS = [
    "Judge novelty annotation",
    "",
    "1. Each row is one pair of research ideas. Read idea_A and idea_B (the two rightmost columns).",
    "2. In 'your_choice' pick A if idea A is more novel, B if idea B is more novel, or unclear if",
    "   the pair cannot be judged.",
    "3. Whenever you pick unclear, write why in 'notes'.",
]


@dataclass
class GroupConfig:
    """One stratum of a batch: where its instances come from and how many to draw."""

    name: str
    run_dir: Path
    per_annotator: int
    select: str = MISTAKES
    # How decisively the judge had to score the pair, as the gap between the two
    # novelty scores on the 0-6 scale. At 6 a control is a 6-vs-0 call and a
    # mistake is confidently wrong rather than an abstention.
    min_novelty_margin: float = 0.0
    test_instances: str | None = None


@dataclass
class BatchConfig:
    out_dir: Path
    num_annotators: int
    groups: list[GroupConfig] = field(default_factory=list)
    assignment: str = DISJOINT
    seed: int = 13
    annotator_names: list[str] | None = None

    @property
    def names(self) -> list[str]:
        if self.annotator_names:
            return [str(n) for n in self.annotator_names]
        return [f"annotator_{i:02d}" for i in range(1, self.num_annotators + 1)]

    @property
    def per_annotator(self) -> int:
        return sum(g.per_annotator for g in self.groups)


def _build(cls, raw: dict, where: str, **overrides):
    unknown = set(raw) - set(cls.__dataclass_fields__)
    if unknown:
        raise SystemExit(f"Unknown key(s) in {where}: {', '.join(sorted(unknown))}")
    try:
        return cls(**{**raw, **overrides})
    except (KeyError, TypeError) as e:
        raise SystemExit(f"Bad {where}: {e}")


def load_config(path: Path) -> BatchConfig:
    raw = yaml.safe_load(path.read_text()) or {}
    groups_raw = raw.pop("groups", None) or []
    if not groups_raw:
        raise SystemExit(f"{path} needs at least one entry under 'groups'.")
    groups = [
        _build(GroupConfig, g, f"group '{g.get('name', '?')}' in {path}", run_dir=Path(g["run_dir"]))
        for g in groups_raw
    ]
    cfg = _build(BatchConfig, raw, str(path), out_dir=Path(raw["out_dir"]), groups=groups)

    names = [g.name for g in groups]
    if len(set(names)) != len(names):
        raise SystemExit(f"Group names must be unique, got {names}.")
    for g in groups:
        if g.select not in (MISTAKES, CONTROLS):
            raise SystemExit(f"Group '{g.name}': select must be '{MISTAKES}' or '{CONTROLS}', not {g.select!r}.")
        if g.per_annotator < 0:
            raise SystemExit(f"Group '{g.name}': per_annotator cannot be negative.")
    if cfg.annotator_names and len(cfg.annotator_names) != cfg.num_annotators:
        raise SystemExit(
            f"annotator_names has {len(cfg.annotator_names)} entries but num_annotators is {cfg.num_annotators}."
        )
    if cfg.num_annotators < 1 or cfg.per_annotator < 1:
        raise SystemExit("num_annotators and the total instances per annotator must be positive.")
    if cfg.assignment not in (DISJOINT, SHARED):
        raise SystemExit(f"assignment must be '{DISJOINT}' or '{SHARED}', not {cfg.assignment!r}.")
    return cfg


def group_pool(group: GroupConfig, repo_root: Path) -> list[ScoredInstance]:
    """The instances a group may draw from: one side of its run's partition, filtered
    down to the pairs the judge was decisive enough about."""
    part = partition_scored(group.run_dir, group.test_instances, repo_root)
    pool = part.mistakes if group.select == MISTAKES else part.controls
    return [s for s in pool if novelty_margin(s[2]) >= group.min_novelty_margin]


def blind_row(scored: ScoredInstance, group: GroupConfig, annotator: str, rng: random.Random) -> dict:
    """One assignment: the blinded view the annotator sees plus its decoding key."""
    pid, inst, sc = scored
    # Idea keys can load from YAML as ints; keep them as str so the key file
    # compares like-typed values against the annotator's decoded answer.
    ideas = {str(k): v for k, v in inst["ideas"].items()}
    a_idea, b_idea = rng.sample(sorted(ideas), 2)
    gold_list = [str(w) for w in inst["expected_winners"]]
    return {
        "annotator": annotator,
        "group": group.name,
        "source": group.run_dir.name,
        "problem_id": pid,
        "area": inst.get("context", ""),
        "A_text": ideas[a_idea],
        "B_text": ideas[b_idea],
        "is_mistake": is_judge_mistake(sc, gold_list),
        "shown_A_idea": a_idea,
        "shown_B_idea": b_idea,
        "gold_idea": gold_list[0],
        "judge_novelty_winner": judge_novelty_winner(sc["novelty"]),
    }


def _deal(pool: list[ScoredInstance], group: GroupConfig, cfg: BatchConfig, rng: random.Random) -> list[ScoredInstance]:
    """Shuffle a group's pool and take exactly as many instances as the batches will need."""
    shared = cfg.assignment == SHARED
    need = group.per_annotator if shared else cfg.num_annotators * group.per_annotator
    if len(pool) < need:
        how = (
            f"one set of {group.per_annotator} shared by {cfg.num_annotators} annotators"
            if shared
            else f"{cfg.num_annotators} annotators x {group.per_annotator} each, disjoint"
        )
        raise SystemExit(f"Group '{group.name}': need {need} instances ({how}) but the pool holds {len(pool)}.")
    return rng.sample(pool, need)


def _reject_duplicates(annotator: str, rows: list[dict]) -> None:
    """Two groups drawing on one run must not hand the same pair over twice."""
    dupes = [k for k, n in Counter((r["source"], r["problem_id"]) for r in rows).items() if n > 1]
    if dupes:
        listed = ", ".join(f"{src}#{pid}" for src, pid in dupes[:5])
        more = f" and {len(dupes) - 5} more" if len(dupes) > 5 else ""
        raise SystemExit(
            f"{annotator}'s batch draws {len(dupes)} instance(s) more than once ({listed}{more}); groups overlap."
        )


def assign_batches(pools: dict[str, list[ScoredInstance]], cfg: BatchConfig) -> dict[str, list[dict]]:
    """Deal each group's instances to the annotators.

    Disjoint gives every annotator different instances, covering more of the
    pools. Shared gives them all the same instances with the same A/B blinding,
    so their answers line up and inter-annotator agreement falls out of a direct
    comparison. Row order is shuffled per annotator either way, so neither
    position nor item_id carries meaning.
    """
    rng = random.Random(cfg.seed)
    dealt = {g.name: _deal(pools[g.name], g, cfg, rng) for g in cfg.groups}
    # Blind the shared sets once, so every annotator sees the same idea as A.
    # Built only when used, to leave the disjoint draws untouched.
    shared = (
        {g.name: [blind_row(s, g, "", rng) for s in dealt[g.name]] for g in cfg.groups}
        if cfg.assignment == SHARED
        else {}
    )

    batches: dict[str, list[dict]] = {}
    for i, name in enumerate(cfg.names):
        rows: list[dict] = []
        for g in cfg.groups:
            if cfg.assignment == SHARED:
                rows += [dict(row, annotator=name) for row in shared[g.name]]
            else:
                rows += [
                    blind_row(s, g, name, rng)
                    for s in dealt[g.name][i * g.per_annotator:(i + 1) * g.per_annotator]
                ]
        # Interleave the groups so nothing about position gives them away, then
        # number the rows: the id is a handle for the annotator, nothing more.
        rng.shuffle(rows)
        for n, row in enumerate(rows, start=1):
            row["item_id"] = f"{n:03d}"
        _reject_duplicates(name, rows)
        batches[name] = rows
    return batches


def write_annotator_xlsx(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "annotation"
    ws.append(SHEET_HEADERS)
    for r in rows:
        ws.append([r["item_id"], r["area"], "", "", r["A_text"], r["B_text"]])

    # Note: do not pass showDropDown -- in the file format that flag *suppresses*
    # the in-cell arrow, so leaving it unset is what makes the dropdown appear.
    dv = DataValidation(type="list", formula1=f'"{",".join(CHOICES)}"', allow_blank=True)
    dv.errorTitle, dv.error = "Invalid choice", f"Pick one of: {', '.join(CHOICES)}."
    ws.add_data_validation(dv)
    dv.add(f"C2:C{len(rows) + 1}")

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, RULED
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    # Wrap the prose columns and top-align every row; heights are left unset so
    # the viewer auto-fits them and no idea text is clipped out of sight.
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=len(rows) + 1)):
        for cell in row:
            cell.border = RULED
            cell.alignment = Alignment(wrap_text=cell.column_letter in WRAPPED_COLUMNS, vertical="top")
            if cell.column_letter == "C":
                # The one cell to act in: tinted, centred and bold so it reads
                # as an input rather than as more text.
                cell.fill = INPUT_FILL
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif i % 2:
                cell.fill = BAND_FILL

    ws.freeze_panes = "E2"
    ws.sheet_view.showGridLines = False  # the ruled borders carry the structure
    ws.sheet_properties.tabColor = ACCENT

    notes = wb.create_sheet("instructions")
    for line in INSTRUCTIONS:
        notes.append([line])
    notes.column_dimensions["A"].width = 100
    for cell in notes["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    notes["A1"].font = Font(bold=True, size=14, color=ACCENT)
    notes.sheet_view.showGridLines = False
    wb.save(path)


def write_key_xlsx(path: Path, batches: dict[str, list[dict]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "key"
    ws.append(KEY_HEADERS)
    for rows in batches.values():
        for r in rows:
            ws.append([r[h] for h in KEY_HEADERS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--config", required=True, help="YAML batch config")
    args = ap.parse_args()

    cfg = load_config(Path(args.config))
    repo_root = Path(__file__).resolve().parents[3]

    pools = {}
    for g in cfg.groups:
        pools[g.name] = group_pool(g, repo_root)
        margin = f", novelty margin >= {g.min_novelty_margin:g}" if g.min_novelty_margin else ""
        print(f"Group '{g.name}': {len(pools[g.name])} available ({g.select}{margin}) in {g.run_dir.name}", flush=True)

    batches = assign_batches(pools, cfg)
    cfg.out_dir.mkdir(parents=True, exist_ok=True)
    mix = " + ".join(f"{g.per_annotator} {g.name}" for g in cfg.groups if g.per_annotator)
    same = " (same set for everyone)" if cfg.assignment == SHARED else ""
    for name, rows in batches.items():
        out = cfg.out_dir / f"{name}.xlsx"
        write_annotator_xlsx(out, rows)
        write_annotation_page(out.with_suffix(".html"), name, rows, CHOICES, INSTRUCTIONS)
        # A third route to the same batch: an Apps Script the batch owner runs
        # once to build this annotator's Google Form, for annotating on a phone.
        script = out.with_name(f"{name}_form.gs")
        script.write_text(
            # One scrolling page, like the CLI's default: with page breaks Submit
            # is only reachable after tapping Next through every pair.
            build_apps_script(name, rows, CHOICES, INSTRUCTIONS, pairs_per_page=0),
            encoding="utf-8",
        )
        print(f"  {out}  (+ .html, + _form.gs) ({mix}){same}")

    key_path = cfg.out_dir / "key.xlsx"
    write_key_xlsx(key_path, batches)
    print(f"Key (do not distribute): {key_path}")


if __name__ == "__main__":
    main()
